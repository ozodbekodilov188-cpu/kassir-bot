import asyncio
import html
import os
import re
import sqlite3
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandObject
from aiogram.types import CallbackQuery, FSInputFile, InlineKeyboardButton, InlineKeyboardMarkup, Message
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = Path(__file__).resolve().parent
ENV_PATH = BASE_DIR / ".env"
load_dotenv(dotenv_path=ENV_PATH, encoding="utf-8-sig", override=True)

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
DB_PATH = str(BASE_DIR / os.getenv("DB_PATH", "kassa.db"))
TZ = ZoneInfo(os.getenv("TIMEZONE", "Asia/Tashkent").strip() or "Asia/Tashkent")

ENV_CASHIERS = {
    int(x.strip()) for x in os.getenv("CASHIER_IDS", "").split(",")
    if x.strip().isdigit()
}
ALLOWED_GROUP_IDS = {
    int(x.strip()) for x in os.getenv("ALLOWED_GROUP_IDS", "").split(",")
    if x.strip().lstrip("-").isdigit()
}

if not BOT_TOKEN:
    raise RuntimeError("BOT_TOKEN topilmadi.")

router = Router()


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def column_exists(conn, table, column):
    cols = conn.execute(f"PRAGMA table_info({table})").fetchall()
    return any(c[1] == column for c in cols)


def init_db():
    with db() as conn:
        conn.execute("""
        CREATE TABLE IF NOT EXISTS payments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            created_at TEXT NOT NULL,
            business_date TEXT NOT NULL,
            group_chat_id INTEGER NOT NULL,
            group_title TEXT,
            group_message_id INTEGER,
            bot_group_message_id INTEGER,
            agent_id INTEGER NOT NULL,
            agent_name TEXT,
            agent_username TEXT,
            client TEXT NOT NULL,
            amount INTEGER NOT NULL,
            comment TEXT,
            status TEXT NOT NULL DEFAULT 'pending',
            cashier_id INTEGER,
            cashier_name TEXT,
            processed_at TEXT
        )
        """)
        conn.execute("""
        CREATE TABLE IF NOT EXISTS cashiers (
            user_id INTEGER PRIMARY KEY,
            name TEXT,
            added_at TEXT NOT NULL
        )
        """)
        conn.execute("""
        CREATE TABLE IF NOT EXISTS admins (
            user_id INTEGER PRIMARY KEY,
            name TEXT,
            added_at TEXT NOT NULL
        )
        """)

        if not column_exists(conn, "payments", "currency"):
            conn.execute("ALTER TABLE payments ADD COLUMN currency TEXT DEFAULT 'UZS'")
        if not column_exists(conn, "payments", "rate"):
            conn.execute("ALTER TABLE payments ADD COLUMN rate INTEGER")
        if not column_exists(conn, "payments", "raw_text"):
            conn.execute("ALTER TABLE payments ADD COLUMN raw_text TEXT")


def now():
    return datetime.now(TZ)


def money(n):
    return f"{int(n):,}".replace(",", " ")


def user_name(user):
    if not user:
        return "Noma'lum"
    name = " ".join(x for x in [user.first_name, user.last_name] if x)
    return name or user.username or str(user.id)


def db_cashiers():
    with db() as conn:
        return {int(r["user_id"]) for r in conn.execute("SELECT user_id FROM cashiers").fetchall()}


def db_admins():
    with db() as conn:
        return {int(r["user_id"]) for r in conn.execute("SELECT user_id FROM admins").fetchall()}


def all_cashiers():
    return ENV_CASHIERS | db_cashiers()


def all_admins():
    return db_admins()


def is_cashier(uid):
    return uid in all_cashiers()


def is_admin(uid):
    return uid in all_admins()


def allowed_group(chat_id):
    return not ALLOWED_GROUP_IDS or chat_id in ALLOWED_GROUP_IDS


def status_text(status):
    return {
        "pending": "⏳ Kutilmoqda",
        "accepted": "✅ Olindi",
        "rejected": "❌ Olinmadi",
    }.get(status, status)


def keyboard(pid):
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="✅ Olindi", callback_data=f"cash:accept:{pid}"),
        InlineKeyboardButton(text="❌ Olinmadi", callback_data=f"cash:reject:{pid}")
    ]])


def payment_text(row):
    rate = money(row["rate"]) if row["rate"] else "—"
    currency = row["currency"] or "UZS"
    return (
        f"💵 <b>Tushum #{row['id']}</b>\n\n"
        f"💱 <b>VALYUTA:</b> {html.escape(currency)}\n"
        f"💰 <b>SUMMA:</b> {money(row['amount'])}\n"
        f"📈 <b>KURS:</b> {rate}\n"
        f"🧑‍💼 <b>AGENT:</b> {html.escape(row['agent_name'] or '')}\n"
        f"👤 <b>KLIENT:</b> {html.escape(row['client'] or '')}\n\n"
        f"📌 <b>HOLAT:</b> {status_text(row['status'])}"
    )


def parse_free_text(text: str):
    original = (text or "").strip()
    if not original:
        return None

    low = original.lower()

    for ch in ("’", "ʻ", "`", "‘"):
        low = low.replace(ch, "'")

    trigger_re = re.compile(
        r"\b(?:tushdi|tushum|tushgan|berdi|keldi|oldik|olindi|"
        r"to'ladi|toladi|o'tkazdi|otkazdi|o'tdi|otdi)\b",
        re.IGNORECASE,
    )

    trigger_match = trigger_re.search(low)
    if not trigger_match:
        return None

    if re.search(r"(?:\$|\busd\b|\bdollar(?:i)?\b|\bdoll\b)", low):
        currency = "USD"
    else:
        currency = "UZS"

    rate = None
    m_rate = re.search(
        r"\b(?:kurs|kursi)\s*[:=\-]?\s*([0-9][0-9\s.,_]*)",
        low
    )
    if m_rate:
        digits = re.sub(r"\D", "", m_rate.group(1))
        if digits:
            rate = int(digits)

    amount = None
    amount_span = None

    mult_patterns = [
        (r"(?<!\w)(\d+(?:[.,]\d+)?)\s*(?:mln|million)\b", 1_000_000),
        (r"(?<!\w)(\d+(?:[.,]\d+)?)\s*(?:ming)\b", 1_000),
        (r"(?<!\w)(\d+(?:[.,]\d+)?)\s*k\b", 1_000),
    ]

    for pattern, multiplier in mult_patterns:
        m = re.search(pattern, low)
        if m:
            try:
                num = float(m.group(1).replace(",", "."))
                amount = int(round(num * multiplier))
                amount_span = m.span()
                break
            except ValueError:
                pass

    if amount is None:
        m = re.search(r"\$\s*([0-9][0-9\s.,_]*)", low)
        if m:
            digits = re.sub(r"\D", "", m.group(1))
            if digits:
                amount = int(digits)
                amount_span = m.span()

    if amount is None:
        prefix = low[:trigger_match.start()]
        candidates = list(
            re.finditer(
                r"(?<!\w)(\d[\d\s.,_]*\d|\d)(?!\w)",
                prefix
            )
        )

        if candidates:
            m = candidates[-1]
            digits = re.sub(r"\D", "", m.group(1))
            if digits:
                amount = int(digits)
                amount_span = m.span()

    if amount is None or amount <= 0 or amount_span is None:
        return None

    before = original[:amount_span[0]].strip(" ,.;:-")

    before = re.sub(
        r"(?i)^\s*(?:tushum|tushdi)\s*[:=\-]?\s*",
        "",
        before
    ).strip()

    before = re.sub(r"(?i)\bakadan\b\s*$", "aka", before).strip()
    before = re.sub(r"(?i)\bopadan\b\s*$", "opa", before).strip()
    before = re.sub(r"(?i)\bukadan\b\s*$", "uka", before).strip()

    before = re.sub(
        r"(?i)\b(?:xorazm|toshkent|samarqand|buxoro|navoiy|namangan|andijon|"
        r"farg'ona|fargona|jizzax|sirdaryo|surxondaryo|qashqadaryo|"
        r"qoraqalpog'iston|qoraqalpogiston)dan\b\s*$",
        "",
        before,
    ).strip(" ,.;:-")

    before = re.sub(
        r"(?i)\b([\w'\-]+)dan\b\s*$",
        r"\1",
        before
    ).strip(" ,.;:-")

    client = re.sub(r"\s+", " ", before).strip() or "Noma'lum"

    if len(client) > 80:
        client = client[:80]

    return {
        "currency": currency,
        "amount": amount,
        "rate": rate,
        "client": client,
        "raw_text": original,
    }


def target_id_from_message(message: Message, command: CommandObject):
    if message.reply_to_message and message.reply_to_message.from_user:
        return (
            message.reply_to_message.from_user.id,
            user_name(message.reply_to_message.from_user)
        )

    if command.args:
        raw = command.args.strip().split()[0]
        if raw.isdigit():
            return int(raw), raw

    return None, None


async def save_and_send_payment(
    message: Message,
    bot: Bot,
    amount: int,
    client: str,
    currency="UZS",
    rate=None,
    comment="",
    raw_text=None
):
    dt = now()

    with db() as conn:
        cur = conn.execute("""
            INSERT INTO payments (
                created_at,
                business_date,
                group_chat_id,
                group_title,
                group_message_id,
                agent_id,
                agent_name,
                agent_username,
                client,
                amount,
                comment,
                currency,
                rate,
                raw_text
            )
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, (
            dt.strftime("%Y-%m-%d %H:%M:%S"),
            dt.strftime("%Y-%m-%d"),
            message.chat.id,
            message.chat.title or "",
            message.message_id,
            message.from_user.id,
            user_name(message.from_user),
            message.from_user.username,
            client,
            amount,
            comment,
            currency,
            rate,
            raw_text
        ))

        pid = cur.lastrowid

        row = conn.execute(
            "SELECT * FROM payments WHERE id=?",
            (pid,)
        ).fetchone()

    group_confirmation = await message.reply(
        "📨 <b>Kassirga yuborildi</b>\n\n" + payment_text(row)
    )

    with db() as conn:
        conn.execute(
            "UPDATE payments SET bot_group_message_id=? WHERE id=?",
            (group_confirmation.message_id, pid)
        )

    delivered = 0

    for cashier_id in all_cashiers():
        try:
            await bot.send_message(
                cashier_id,
                "🔔 <b>Yangi tushum keldi</b>\n\n" + payment_text(row),
                reply_markup=keyboard(pid)
            )
            delivered += 1
        except Exception:
            pass

    if delivered == 0:
        await message.reply(
            "⚠️ Kassir botga private chatda /start yuborishi kerak."
        )


@router.message(Command("start"))
async def start(message: Message):
    text = (
        "💰 <b>Kassir Nazorat Bot</b>\n\n"
        "Guruhda oddiy yozishingiz mumkin:\n"
        "<code>Nurali akadan 5000000 tushdi</code>\n"
        "<code>Akmal 25 mln so'm tushdi</code>\n"
        "<code>Jaxongir aka 4000$ tushdi kurs 12650</code>\n\n"
        "/bugun — bugungi statistika\n"
        "/hisobot — Excel hisobot\n"
        "/id — Telegram ID"
    )

    if not all_admins():
        text += (
            "\n\n👑 Birinchi sozlash: "
            "egasi private chatda /claimadmin yuborsin."
        )

    if is_admin(message.from_user.id):
        text += "\n\n/adminhelp — admin komandalar"

    await message.answer(text)


@router.message(Command("id"))
async def show_id(message: Message):
    await message.answer(
        f"User ID: <code>{message.from_user.id}</code>\n"
        f"Chat ID: <code>{message.chat.id}</code>"
    )


@router.message(Command("claimadmin"))
async def claim_admin(message: Message):
    if message.chat.type != "private":
        await message.answer(
            "Bu komandani botning shaxsiy chatida yuboring."
        )
        return

    if all_admins():
        if is_admin(message.from_user.id):
            await message.answer("✅ Siz allaqachon adminsiz.")
        else:
            await message.answer("❌ Admin allaqachon belgilangan.")
        return

    with db() as conn:
        conn.execute(
            "INSERT OR REPLACE INTO admins(user_id,name,added_at) "
            "VALUES(?,?,?)",
            (
                message.from_user.id,
                user_name(message.from_user),
                now().strftime("%Y-%m-%d %H:%M:%S")
            )
        )

    await message.answer(
        "👑 <b>Siz admin bo'ldingiz.</b>\n\n"
        "/adminhelp — komandalar"
    )


@router.message(Command("adminhelp"))
async def admin_help(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu komanda faqat admin uchun.")
        return

    await message.answer(
        "👑 <b>Admin komandalar</b>\n\n"
        "/cashiers — kassirlar\n"
        "/addcashier USER_ID — kassir qo'shish\n"
        "/removecashier USER_ID — kassirni o'chirish\n"
        "/setcashier USER_ID — kassirni almashtirish"
    )


@router.message(Command("cashiers"))
async def list_cashiers(message: Message):
    if not is_admin(message.from_user.id):
        await message.answer("Bu komanda faqat admin uchun.")
        return

    with db() as conn:
        rows = conn.execute(
            "SELECT user_id,name FROM cashiers ORDER BY added_at"
        ).fetchall()

    lines = [
        f"• {html.escape(r['name'] or '')} — "
        f"<code>{r['user_id']}</code>"
        for r in rows
    ]

    for uid in ENV_CASHIERS:
        if not any(int(r["user_id"]) == uid for r in rows):
            lines.append(
                f"• ENV kassir — <code>{uid}</code>"
            )

    await message.answer(
        "💼 <b>Kassirlar:</b>\n\n"
        + ("\n".join(lines) if lines else "Hozir kassir yo'q.")
    )


@router.message(Command("addcashier"))
async def add_cashier(
    message: Message,
    command: CommandObject
):
    if not is_admin(message.from_user.id):
        await message.answer("Bu komanda faqat admin uchun.")
        return

    uid, name = target_id_from_message(message, command)

    if not uid:
        await message.answer(
            "<code>/addcashier 123456789</code>"
        )
        return

    with db() as conn:
        conn.execute(
            "INSERT OR REPLACE INTO cashiers(user_id,name,added_at) "
            "VALUES(?,?,?)",
            (
                uid,
                name or str(uid),
                now().strftime("%Y-%m-%d %H:%M:%S")
            )
        )

    await message.answer(
        f"✅ Kassir qo'shildi: <code>{uid}</code>"
    )


@router.message(Command("removecashier"))
async def remove_cashier(
    message: Message,
    command: CommandObject
):
    if not is_admin(message.from_user.id):
        await message.answer("Bu komanda faqat admin uchun.")
        return

    uid, _ = target_id_from_message(message, command)

    if not uid:
        await message.answer(
            "<code>/removecashier 123456789</code>"
        )
        return

    if uid in ENV_CASHIERS:
        await message.answer(
            "⚠️ Bu kassir .env orqali qo'shilgan."
        )
        return

    with db() as conn:
        cur = conn.execute(
            "DELETE FROM cashiers WHERE user_id=?",
            (uid,)
        )

    await message.answer(
        "🗑 Kassir o'chirildi."
        if cur.rowcount
        else "Bunday kassir topilmadi."
    )


@router.message(Command("setcashier"))
async def set_cashier(
    message: Message,
    command: CommandObject
):
    if not is_admin(message.from_user.id):
        await message.answer("Bu komanda faqat admin uchun.")
        return

    uid, name = target_id_from_message(message, command)

    if not uid:
        await message.answer(
            "<code>/setcashier 123456789</code>"
        )
        return

    with db() as conn:
        conn.execute("DELETE FROM cashiers")

        conn.execute(
            "INSERT OR REPLACE INTO cashiers(user_id,name,added_at) "
            "VALUES(?,?,?)",
            (
                uid,
                name or str(uid),
                now().strftime("%Y-%m-%d %H:%M:%S")
            )
        )

    await message.answer(
        f"🔄 Yangi kassir o'rnatildi: <code>{uid}</code>"
    )


@router.message(Command("claimcashier"))
async def claim_cashier_legacy(message: Message):
    await message.answer(
        "Kassirni Admin /addcashier yoki "
        "/setcashier orqali belgilaydi."
    )


@router.message(Command("tushum"))
async def add_payment_command(
    message: Message,
    command: CommandObject,
    bot: Bot
):
    if message.chat.type not in ("group", "supergroup"):
        await message.answer(
            "Tushumni guruh ichida kiriting."
        )
        return

    if not allowed_group(message.chat.id):
        return

    if not all_cashiers():
        await message.answer(
            "⚠️ Admin avval kassirni belgilashi kerak."
        )
        return

    if not command.args:
        await message.answer(
            "Misol: "
            "<code>/tushum 250000 MijozNomi izoh</code>"
        )
        return

    parts = command.args.strip().split(maxsplit=2)

    if len(parts) < 2:
        await message.answer(
            "Summa va klient nomini kiriting."
        )
        return

    raw = re.sub(r"\D", "", parts[0])

    if not raw:
        await message.answer("Summa noto'g'ri.")
        return

    amount = int(raw)
    client = parts[1]
    comment = parts[2] if len(parts) > 2 else ""

    await save_and_send_payment(
        message,
        bot,
        amount,
        client,
        "UZS",
        None,
        comment,
        command.args
    )


@router.message(F.text & ~F.text.startswith("/"))
async def free_text_payment(
    message: Message,
    bot: Bot
):
    if message.chat.type not in ("group", "supergroup"):
        return

    if not allowed_group(message.chat.id):
        return

    if not all_cashiers():
        return

    data = parse_free_text(message.text)

    if not data:
        return

    await save_and_send_payment(
        message,
        bot,
        amount=data["amount"],
        client=data["client"],
        currency=data["currency"],
        rate=data["rate"],
        comment="",
        raw_text=data["raw_text"]
    )


@router.callback_query(F.data.startswith("cash:"))
async def cashier_action(
    callback: CallbackQuery,
    bot: Bot
):
    if not is_cashier(callback.from_user.id):
        await callback.answer(
            "Bu tugma faqat kassir uchun.",
            show_alert=True
        )
        return

    _, action, pid = callback.data.split(":")
    pid = int(pid)

    new_status = (
        "accepted"
        if action == "accept"
        else "rejected"
    )

    with db() as conn:
        row = conn.execute(
            "SELECT * FROM payments WHERE id=?",
            (pid,)
        ).fetchone()

        if not row:
            await callback.answer(
                "Tushum topilmadi.",
                show_alert=True
            )
            return

        if row["status"] != "pending":
            await callback.answer(
                "Bu tushum allaqachon tekshirilgan.",
                show_alert=True
            )
            return

        conn.execute("""
            UPDATE payments
            SET status=?,
                cashier_id=?,
                cashier_name=?,
                processed_at=?
            WHERE id=?
        """, (
            new_status,
            callback.from_user.id,
            user_name(callback.from_user),
            now().strftime("%Y-%m-%d %H:%M:%S"),
            pid
        ))

        row = conn.execute(
            "SELECT * FROM payments WHERE id=?",
            (pid,)
        ).fetchone()

    try:
        await callback.message.edit_text(
            "✅ <b>Natija saqlandi</b>\n\n"
            + payment_text(row)
        )
    except Exception:
        pass

    result_word = (
        "✅ OLINDI"
        if new_status == "accepted"
        else "❌ OLINMADI"
    )

    if row["bot_group_message_id"]:
        try:
            await bot.edit_message_text(
                chat_id=row["group_chat_id"],
                message_id=row["bot_group_message_id"],
                text=f"{result_word}\n\n{payment_text(row)}"
            )
        except Exception:
            pass

    try:
        await bot.send_message(
            row["agent_id"],
            f"{result_word}\n\n"
            f"💱 {row['currency'] or 'UZS'}\n"
            f"💰 {money(row['amount'])}\n"
            f"👤 {html.escape(row['client'])}"
        )
    except Exception:
        pass

    await callback.answer(
        "Olindi ✅"
        if new_status == "accepted"
        else "Olinmadi ❌"
    )


@router.message(Command("bugun"))
async def today(message: Message):
    if (
        not is_cashier(message.from_user.id)
        and not is_admin(message.from_user.id)
    ):
        await message.answer(
            "Bu komanda kassir yoki admin uchun."
        )
        return

    d = now().strftime("%Y-%m-%d")

    with db() as conn:
        rows = conn.execute("""
            SELECT
                currency,
                status,
                COUNT(*) cnt,
                COALESCE(SUM(amount),0) total
            FROM payments
            WHERE business_date=?
            GROUP BY currency,status
            ORDER BY currency,status
        """, (d,)).fetchall()

    if not rows:
        await message.answer(
            f"📊 {d}\nBugun tushum yo'q."
        )
        return

    lines = [f"📊 <b>{d}</b>"]

    for r in rows:
        lines.append(
            f"{r['currency'] or 'UZS'} | "
            f"{status_text(r['status'])}: "
            f"{r['cnt']} ta — "
            f"<b>{money(r['total'])}</b>"
        )

    await message.answer("\n".join(lines))


def make_excel(report_date):
    with db() as conn:
        rows = conn.execute(
            "SELECT * FROM payments "
            "WHERE business_date=? "
            "ORDER BY created_at",
            (report_date,)
        ).fetchall()

    if not rows:
        return None

    out_dir = BASE_DIR / "reports"
    out_dir.mkdir(exist_ok=True)

    path = out_dir / f"kassa_{report_date}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Kunlik hisobot"

    dark_blue = "1F4E78"
    light_blue = "D9EAF7"
    light_green = "E2F0D9"
    light_red = "FCE4D6"
    light_yellow = "FFF2CC"
    white = "FFFFFF"
    border_color = "D9E1F2"

    thin = Side(
        style="thin",
        color=border_color
    )

    all_border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    ws.merge_cells("A1:K1")

    ws["A1"] = (
        f"KASSA KUNLIK HISOBOTI — {report_date}"
    )

    ws["A1"].font = Font(
        size=16,
        bold=True,
        color=white
    )

    ws["A1"].fill = PatternFill(
        "solid",
        fgColor=dark_blue
    )

    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws.row_dimensions[1].height = 28

    totals = {}

    for row in rows:
        cur = row["currency"] or "UZS"
        status = row["status"]

        totals.setdefault(
            cur,
            {
                "accepted": 0,
                "pending": 0,
                "rejected": 0
            }
        )

        totals[cur][status] = (
            totals[cur].get(status, 0)
            + int(row["amount"])
        )

    ws["A3"] = "KUNLIK YAKUN"
    ws["A3"].font = Font(
        bold=True,
        color=dark_blue
    )

    summary_row = 4

    for cur, vals in sorted(totals.items()):
        ws.cell(summary_row, 1, cur)
        ws.cell(summary_row, 2, "Olindi")
        ws.cell(
            summary_row,
            3,
            vals.get("accepted", 0)
        )

        ws.cell(summary_row, 4, "Kutilmoqda")
        ws.cell(
            summary_row,
            5,
            vals.get("pending", 0)
        )

        ws.cell(summary_row, 6, "Olinmadi")
        ws.cell(
            summary_row,
            7,
            vals.get("rejected", 0)
        )

        for c in range(1, 8):
            cell = ws.cell(summary_row, c)
            cell.border = all_border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        ws.cell(
            summary_row,
            1
        ).font = Font(bold=True)

        ws.cell(
            summary_row,
            2
        ).fill = PatternFill(
            "solid",
            fgColor=light_green
        )

        ws.cell(
            summary_row,
            4
        ).fill = PatternFill(
            "solid",
            fgColor=light_yellow
        )

        ws.cell(
            summary_row,
            6
        ).fill = PatternFill(
            "solid",
            fgColor=light_red
        )

        ws.cell(
            summary_row,
            3
        ).number_format = "#,##0"

        ws.cell(
            summary_row,
            5
        ).number_format = "#,##0"

        ws.cell(
            summary_row,
            7
        ).number_format = "#,##0"

        summary_row += 1

    header_row = summary_row + 2

    headers = [
        "№",
        "VAQT",
        "VALYUTA",
        "SUMMA",
        "KURS",
        "AGENT",
        "KLIENT",
        "HOLAT",
        "KASSIR",
        "GURUH",
        "ASL XABAR"
    ]

    for col, title in enumerate(headers, 1):
        cell = ws.cell(header_row, col, title)

        cell.font = Font(
            bold=True,
            color=white
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=dark_blue
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = all_border

    ws.row_dimensions[header_row].height = 24

    for i, row in enumerate(rows, 1):
        excel_row = header_row + i

        status = (
            status_text(row["status"])
            .replace("✅ ", "")
            .replace("❌ ", "")
            .replace("⏳ ", "")
        )

        values = [
            i,
            row["created_at"],
            row["currency"] or "UZS",
            row["amount"],
            row["rate"] or "",
            row["agent_name"] or "",
            row["client"] or "",
            status,
            row["cashier_name"] or "",
            row["group_title"] or "",
            row["raw_text"] or "",
        ]

        for col, value in enumerate(values, 1):
            cell = ws.cell(
                excel_row,
                col,
                value
            )

            cell.border = all_border

            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

        ws.cell(
            excel_row,
            4
        ).number_format = "#,##0"

        ws.cell(
            excel_row,
            5
        ).number_format = "#,##0"

        status_cell = ws.cell(
            excel_row,
            8
        )

        status_cell.font = Font(
            bold=True
        )

        status_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        if row["status"] == "accepted":
            status_cell.fill = PatternFill(
                "solid",
                fgColor=light_green
            )

        elif row["status"] == "rejected":
            status_cell.fill = PatternFill(
                "solid",
                fgColor=light_red
            )

        else:
            status_cell.fill = PatternFill(
                "solid",
                fgColor=light_yellow
            )

        if (row["currency"] or "UZS") == "USD":
            ws.cell(
                excel_row,
                3
            ).fill = PatternFill(
                "solid",
                fgColor=light_blue
            )

            ws.cell(
                excel_row,
                3
            ).font = Font(
                bold=True
            )

    last_row = header_row + len(rows)

    ws.auto_filter.ref = (
        f"A{header_row}:K{last_row}"
    )

    ws.freeze_panes = (
        f"A{header_row + 1}"
    )

    widths = {
        "A": 6,
        "B": 20,
        "C": 12,
        "D": 16,
        "E": 14,
        "F": 22,
        "G": 28,
        "H": 16,
        "I": 22,
        "J": 24,
        "K": 42
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    total_start = last_row + 3

    ws.cell(
        total_start,
        1,
        "JAMI OLINGAN"
    )

    ws.cell(
        total_start,
        1
    ).font = Font(
        bold=True,
        color=dark_blue
    )

    rr = total_start + 1

    for cur, vals in sorted(totals.items()):
        ws.cell(rr, 1, cur)

        ws.cell(
            rr,
            2,
            vals.get("accepted", 0)
        )

        ws.cell(
            rr,
            1
        ).font = Font(
            bold=True
        )

        ws.cell(
            rr,
            2
        ).font = Font(
            bold=True
        )

        ws.cell(
            rr,
            2
        ).number_format = "#,##0"

        ws.cell(
            rr,
            1
        ).fill = PatternFill(
            "solid",
            fgColor=light_green
        )

        ws.cell(
            rr,
            2
        ).fill = PatternFill(
            "solid",
            fgColor=light_green
        )

        ws.cell(
            rr,
            1
        ).border = all_border

        ws.cell(
            rr,
            2
        ).border = all_border

        rr += 1

    wb.save(path)

    return path


@router.message(Command("hisobot"))
async def report(
    message: Message,
    command: CommandObject
):
    if (
        not is_cashier(message.from_user.id)
        and not is_admin(message.from_user.id)
    ):
        await message.answer(
            "Bu komanda kassir yoki admin uchun."
        )
        return

    d = now().strftime("%Y-%m-%d")

    if command.args:
        try:
            datetime.strptime(
                command.args.strip(),
                "%Y-%m-%d"
            )

            d = command.args.strip()

        except ValueError:
            await message.answer(
                "Format: /hisobot 2026-09-12"
            )
            return

    path = make_excel(d)

    if not path:
        await message.answer(
            "Bu sana uchun ma'lumot yo'q."
        )
        return

    await message.answer_document(
        FSInputFile(path),
        caption=f"📄 {d} kassa hisoboti"
    )


async def main():
    init_db()

    bot = Bot(
        BOT_TOKEN,
        default=DefaultBotProperties(
            parse_mode=ParseMode.HTML
        )
    )

    dp = Dispatcher()
    dp.include_router(router)

    print("BOT ISHLAYAPTI. Bu oynani yopmang.")

    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
